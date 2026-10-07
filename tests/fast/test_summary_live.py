"""The district summary's live section, in English and Thai (ADR-0056, summary amendment)."""

from __future__ import annotations

import base64
import io
from datetime import UTC, datetime

import pytest
from docx import Document

import core.flood_evidence.summary_live as summary_live
from core.flood_evidence.planner_answer import NO_WARNINGS
from core.summary_docx import render_summary

NOW = datetime(2026, 10, 4, 7, 20, tzinfo=UTC)
PICTURE = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
FACTS = [
    {"id": "S", "kind": "situation", "active_incidents": 2, "conflicting_incidents": 0,
     "incidents_by_confidence": {"conflicting": 0, "low": 1, "medium": 1, "high": 0}},
    {"id": "C", "kind": "changes", "window_minutes": 60, "new_incidents": 1,
     "incidents_grew": 0, "no_longer_reported": 1},
    {"id": "I1", "kind": "incident", "roads": ["Chalong Krung Road"], "confidence": "low",
     "reports": 2, "deepest_reported_cm": 20, "newest_evidence_bangkok": "14:05"},
    {"id": "I2", "kind": "incident", "roads": ["Luang Phaeng Road"], "confidence": "medium",
     "reports": 3, "deepest_reported_cm": None, "newest_evidence_bangkok": "14:10"},
    {"id": "F1", "kind": "facility", "name": "วัดหนองจอก", "facility_type": "evacuation_centre",
     "flooding_reported_within_m": 90, "source": "DDPM evacuation centre (GRP data library)"},
    {"id": "L", "kind": "limits", "incidents_not_listed": 0},
]
ROAD = {"type": "LineString", "coordinates": [[100.75, 13.72], [100.751, 13.72]]}


def _camera(camera_id: str, provider: str, distance: int) -> dict:
    return {"camera_id": camera_id, "provider": provider, "provider_camera_id": camera_id[-4:],
            "name": f"Camera {camera_id}", "viewer_url": None, "placeholder": False,
            "source_label": "BMA traffic", "distance_m": distance}


@pytest.fixture
def stored(monkeypatch):
    monkeypatch.setattr(summary_live, "live_facts", lambda *_a, **_k: {
        "available": True, "district_code": "1011", "district_name": "Lat Krabang",
        "district_name_th": "ลาดกระบัง", "rolled_up_from": None, "facts": FACTS,
        "labels": {"I1": "inc-1", "I2": "inc-2", "F1": "ddpm:1"}, "as_of": NOW.isoformat()})
    monkeypatch.setattr(summary_live, "list_incidents", lambda *_a: {"incidents": [
        {"incident_id": "inc-1", "road_keys": ["a" * 16]},
        {"incident_id": "inc-2", "road_keys": ["b" * 16]},
    ]})
    monkeypatch.setattr(summary_live, "current_roads", lambda *_a, **_k: {
        "features": [{"properties": {"id": k * 16}, "geometry": ROAD} for k in "ab"],
        "snapshot_retrieved_at": "2026-10-04T07:10:00+00:00"})
    monkeypatch.setattr(summary_live, "camera_registry", lambda _pilot: ())
    monkeypatch.setattr(summary_live, "nearby_cameras", lambda *_a: [
        _camera("bmatraffic:1362", "BMA_TRAFFIC", 120), _camera("itic:9", "ITIC", 50)])


def test_the_live_section_gathers_incidents_facilities_and_one_camera_per_incident(stored) -> None:
    fetched = []

    def frame(camera):
        fetched.append(camera["provider_camera_id"])
        return PICTURE, NOW

    live = summary_live.live_section(None, "adpc", "1011", "district", NOW, get_frame=frame)
    assert live["available"] and live["as_of"] == "04 Oct 2026 14:10"
    assert [i["id"] for i in live["incidents"]] == ["I1", "I2"]
    assert live["facilities"][0]["facility_type"] == "evacuation_centre"
    first, second = live["cameras"]
    assert first["picture"] == PICTURE and first["credit"] == summary_live.CAMERA_CREDIT
    # A camera already used for one incident is not repeated for the next.
    assert second["name"] == "Camera itic:9" and second["picture"] is None
    assert fetched == ["1362"]


def test_a_camera_that_does_not_answer_is_listed_without_a_picture(stored) -> None:
    def broken(_camera):
        raise RuntimeError("no picture")

    live = summary_live.live_section(None, "adpc", "1011", "district", NOW, get_frame=broken)
    assert live["cameras"][0]["picture"] is None
    assert live["cameras"][0]["picture_note"] == "camera did not answer"


def test_camera_pictures_follow_the_deployment_switch(stored) -> None:
    called = False

    def frame(_camera):
        nonlocal called
        called = True
        return PICTURE, NOW

    live = summary_live.live_section(
        None, "adpc", "1011", "district", NOW, get_frame=frame,
        include_camera_pictures=False,
    )
    assert live["cameras"][0]["picture"] is None
    assert not called


def test_municipal_snapshot_camera_can_appear_in_the_word_summary(stored, monkeypatch) -> None:
    camera = _camera("pakkret:9", "PAKKRET_CCTV", 40)
    camera["_snapshot_url"] = "https://example.test/camera.jpg"
    camera["source_label"] = "Pak Kret Municipality CCTV"
    monkeypatch.setattr(summary_live, "nearby_cameras", lambda *_a: [camera])

    live = summary_live.live_section(
        None, "adpc", "1206", "district", NOW,
        get_frame=lambda _camera: (PICTURE, NOW),
    )

    assert live["cameras"][0]["picture"] == PICTURE
    assert live["cameras"][0]["credit"] == "Pak Kret Municipality CCTV"


def _text(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    return "\n".join([p.text for p in document.paragraphs]
                     + [c.text for t in document.tables for r in t.rows for c in r.cells])


def _facts(live: dict) -> dict:
    return {"generated_at": "04 Oct 2026 14:20", "prepared_by": "Planner", "hub": "ADPC Hub",
            "area": {"name": "LAT KRABANG", "name_th": "ลาดกระบัง", "admin_level": "district",
                     "province_name": "BANGKOK"},
            "centres": {"total": 0, "rows": [], "gap_note": "No centres.",
                        "source_line": "Shelters"},
            "live": live, "limits": [], "sources": []}


def test_the_english_summary_has_the_live_section_after_global_risk_with_a_picture(stored) -> None:
    live = summary_live.live_section(None, "adpc", "1011", "district", NOW,
                                     get_frame=lambda _c: (PICTURE, NOW))
    out = render_summary(_facts(live), "en")
    text = _text(out)
    assert text.index("5. Global Risk evidence") < text.index("6. Live reported flooding")
    assert text.index("6. Live reported flooding") < text.index("7. River outlook")
    assert text.index("7. River outlook") < text.index("8. What this cannot tell you")
    assert "As of 04 Oct 2026 14:10, Bangkok time" in text and "not a warning" in text
    assert "DDPM evacuation centre" in text and "Chalong Krung Road" in text
    assert NO_WARNINGS["en"] in text and NO_WARNINGS["th"] in text
    assert "officer" not in text.lower()
    assert len(Document(io.BytesIO(out)).inline_shapes) == 1


def test_the_thai_summary_translates_fixed_text_and_keeps_the_rules(stored) -> None:
    live = summary_live.live_section(None, "adpc", "1011", "district", NOW,
                                     get_frame=lambda _c: (PICTURE, NOW))
    text = _text(render_summary(_facts(live), "th"))
    assert "สรุปการเตรียมความพร้อมรับมือน้ำท่วม" in text
    assert "6. รายงานน้ำท่วมแบบสด" in text and "ศูนย์พักพิง ปภ." in text
    assert NO_WARNINGS["th"] in text
    assert "1. At a glance" not in text


def test_outside_bangkok_the_live_section_is_one_plain_line() -> None:
    text = _text(render_summary(_facts({"available": False, "reason": "outside_coverage"}), "en"))
    assert "6. Live reported flooding" in text and "covers Bangkok and Nonthaburi only" in text


def test_a_camera_without_a_real_name_is_called_by_its_number() -> None:
    assert summary_live._camera_name({"name": " - · - ", "provider_camera_id": "1362",
                                      "camera_id": "bmatraffic:1362"}) == "Camera 1362"
    named = {"en": "On Nut Road", "th": "ถนนอ่อนนุช"}
    assert summary_live._camera_name({"name": named, "camera_id": "x"}) == named


def test_a_blank_name_in_one_language_falls_back_to_the_other() -> None:
    from core.summary_docx import _local

    name = {"en": "- · -", "th": "อุโมงค์ทางลอดพระราม 9"}
    assert _local(name, "en") == "อุโมงค์ทางลอดพระราม 9"
    assert _local({"en": "On Nut Road", "th": "ถนนอ่อนนุช"}, "en") == "On Nut Road"


def test_the_river_outlook_section_shows_trend_peak_chart_and_caveats() -> None:
    from api.river_outlook import chart

    series = [{"valid_at_utc": f"2026-10-0{4 + d}T00:00:00Z", "median_m3s": 900.0 + 60 * d,
               "p25_m3s": 800.0 + 50 * d, "p75_m3s": 1000.0 + 70 * d} for d in range(5)]
    png = chart(series)
    assert png and png[1:4] == b"PNG"
    river = {"available": True, "chart": png, "problem": None,
             "reach": {"reach_id": 430537201, "why": "chao_phraya_nonthaburi",
                       "likely_name_en": "Chao Phraya River", "likely_name_th": "แม่น้ำเจ้าพระยา",
                       "inland": False},
             "summary": {"trend": "rise", "first_median_m3s": 900.0, "median_peak_m3s": 1140.0,
                         "p25_at_peak_m3s": 1000.0, "p75_at_peak_m3s": 1280.0,
                         "median_peak_valid_at_utc": "2026-10-08T00:00:00Z",
                         "issued_at_utc": "2026-10-04T00:00:00Z", "quality_state": "latest",
                         "run": "2026100400"}}
    facts = {**_facts({"available": False, "reason": "outside_coverage"}), "river": river}
    english = render_summary(facts, "en")
    text = _text(english)
    assert "7. River outlook (GEOGLOWS, exploratory)" in text
    assert "forecast to be rising" in text and "1,140 m³/s" in text
    assert "Not confirmed by a hydrologist" in text and "not water level" in text
    assert len(Document(io.BytesIO(english)).inline_shapes) == 1
    thai = _text(render_summary(facts, "th"))
    assert "แนวโน้มแม่น้ำ" in thai and "เพิ่มขึ้น" in thai and "แม่น้ำเจ้าพระยา" in thai


def test_an_area_without_a_reach_says_so_plainly() -> None:
    text = _text(render_summary({**_facts(None), "river": {"available": False,
                                                           "reason": "no_reach"}}, "en"))
    assert "No GEOGLOWS river reach is linked to this area" in text
