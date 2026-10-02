"""ADR-0033: the district summary carries only this district's Global Risk evidence."""

from datetime import datetime
from types import SimpleNamespace

from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from api.planning_summary import _global_risk, global_risk_stats
from core.access_models import AppUser, Base, Hub
from core.assessment_models import Boundary
from core.planning_memory_models import PlanningChatMessage
from grpcli.admin import assign_member, bootstrap_platform_admin, ensure_hub


def _boundary(name: str, code: str, province: str) -> Boundary:
    return Boundary(
        admin_code=code, admin_level="district", name=name, province_name=province,
        country_name="Thailand",
        geom={"type": "Polygon", "coordinates": [[[100, 13], [101, 13], [101, 14], [100, 13]]]},
        source="Thailand hierarchy delivery", edition="2025-10", geometry_sha256="a" * 64,
        is_supported=True,
    )


def _evidence(requested: str, **extra) -> dict:
    return {
        "answer": "## In short\\n3 of 9 schools [1]",
        "answer_source": "draft",
        "map_link": "https://sig.example/?embed=hazard_map&receipt_id=r1",
        "map_link_verified": False,
        "evidence": {
            "pack_id": f"pack-{requested[:4]}",
            "assembled_at": "2026-09-28T03:00:00+00:00",
            "area": {"requested": requested, "sig_place": requested.split(",")[0]},
            "summary": {"sources": 3, "pulled_live": 0, "computed": 3, "declared_gaps": 1},
            "stats": {"hazard": "hazard_flood", "aoi": {"area_km2": 28}, "counts": {
                "schools": {"exposed": 3, "total": 9, "at_risk": 2, "by_risk": {"1": 2}},
                "roads": {"exposed_km": 1.5, "total_km": 10.0},
            }},
            "citations": [{"n": 1, "title": "schools vs flood", "source": "OSM"}],
            "gaps": ["vintages unknown"],
            "receipt": {"receipt_id": "r1", "public_url": "https://sig.example/r/r1"},
            **extra,
        },
    }


def test_only_this_districts_evidence_is_used_and_risk_fields_are_left_out() -> None:
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        bootstrap_platform_admin(session, "owner@example.test")
        ensure_hub(session, actor_email="owner@example.test", code="adpc", name="ADPC Hub")
        assign_member(session, actor_email="owner@example.test", email="planner@example.test",
                      hub_code="adpc", role="planner")
        user = session.scalar(select(AppUser).where(AppUser.email == "planner@example.test"))
        hub = session.scalar(select(Hub).where(Hub.code == "adpc"))
        bang_kapi = _boundary("BANG KAPI", "1006", "BANGKOK")
        session.add(bang_kapi)
        # Bang Phli is newer, so a "latest evidence" shortcut would pick it.
        for requested, when in (
            ("BANG KAPI District, BANGKOK, Thailand", datetime(2026, 9, 28, 3)),
            ("BANG PHLI District, SAMUT PRAKAN, Thailand", datetime(2026, 9, 28, 4)),
        ):
            session.add(PlanningChatMessage(
                user_id=user.id, hub_id=hub.id, role="assistant", kind="evidence",
                text="", question=f"About {requested}", payload=_evidence(requested),
                created_at=when,
            ))
        session.commit()

        found = _global_risk(session, SimpleNamespace(user_id=user.id), hub.id, bang_kapi)

    assert found is not None
    assert found["place"] == "BANG KAPI District"
    assert found["question"] == "About BANG KAPI District, BANGKOK, Thailand"
    assert [row[:4] for row in found["stats"]] == [
        ["Schools", 3, 9, "count"], ["Roads", 1.5, 10.0, "km"],
    ]
    assert "Global Risk adds it" in found["stats"][0][5]
    assert found["hazard"] == "hazard_flood" and found["area_km2"] == 28
    assert "receipt r1" in found["status"]
    assert found["map_link"].endswith("receipt_id=r1")
    assert found["source_note"] and "not a report of current flooding" in found["source_note"]


def test_global_risk_rows_name_items_plainly_and_never_carry_risk_levels() -> None:
    rows = global_risk_stats({"stats": {"counts": {
        "evacuation_centres_th_test": {
            "total": 16, "exposed": 2, "at_risk": 1, "by_risk": {"3": 1},
            "by_severity": {"1": 0, "2": 2},
        },
        "roads": {"total_km": 522.8, "exposed_km": 3.0, "at_risk_km": 2.5},
    }}})

    assert rows[0][:5] == ["Evacuation centres, test upload (evacuation_centres_th_test)", 2, 16,
                           "count", "class 2: 2"]
    # ADPC's test upload of GRP's own centres is not a second opinion.
    assert "not an independent check" in rows[0][5]
    assert rows[1][:4] == ["Roads", 3.0, 522.8, "km"]
    assert 2.5 not in rows[1] and 1 not in rows[0][1:3]


def test_grp_data_counted_back_is_never_called_a_global_risk_addition() -> None:
    counts = {name: {"total": 4, "exposed": 1, "by_severity": {"1": 1}} for name in (
        "evacuation_centres_ddpm", "early_warning_towers_ddpm",
        "civil_defence_volunteer_centres_ddpm", "villages_th_register",
        "our_hub_layer", "evacuation_centres_other_agency",
    )}
    rows = {row[0]: row[5] for row in global_risk_stats(
        {"stats": {"counts": counts}}, frozenset({"our_hub_layer"}),
    )}

    for label in ("Evacuation centres (evacuation_centres_ddpm)",
                  "Early-warning towers (early_warning_towers_ddpm)",
                  "Civil-defence volunteer centres (civil_defence_volunteer_centres_ddpm)",
                  "Villages (villages_th_register)",
                  "Our hub layer (our_hub_layer)"):
        assert "GRP's own data" in rows[label]
    # Another agency's centres are matched by exact name, not by the word "evacuation".
    other = rows["Evacuation centres other agency (evacuation_centres_other_agency)"]
    assert "Not in GRP's data" in other and "GRP's own" not in other


def test_a_count_in_an_unknown_shape_keeps_a_row_pointing_to_the_brief() -> None:
    rows = global_risk_stats({"stats": {"counts": {
        "population_th_village_register": {"people": 12633, "by_severity": {"1": 5000}},
    }}})

    assert rows == [["People, village register grid (population_th_village_register)",
                     "-", "-", "-", "-", rows[0][5]]]
    assert "GRP's own data" in rows[0][5] and "read it in the brief" in rows[0][5]


def test_the_word_summary_names_global_risks_flood_layer_and_polygon() -> None:
    import io

    from docx import Document

    from core.summary_docx import render_summary

    facts = {
        "generated_at": "29 Sep 2026 10:00", "prepared_by": "Planner", "hub": "ADPC Hub",
        "area": {"name": "SAMKO", "province_name": "ANG THONG"},
        "centres": {"total": 0, "rows": [], "gap_note": "No centres.", "source_line": "Shelters"},
        "global_risk": {
            "place": "Samko District", "pack_id": "p1", "assembled_at": "29 Sep 2026 09:00",
            "status": "unverified draft, not published", "hazard": "hazard_flood",
            "area_km2": 1784,
            "stats": global_risk_stats({"stats": {"counts": {
                "villages_th_register": {"total": 33, "exposed": 25, "by_severity": {"1": 9}},
            }}}),
        },
    }
    text = "\n".join(
        [p.text for p in Document(io.BytesIO(render_summary(facts))).paragraphs]
        + [cell.text for table in Document(io.BytesIO(render_summary(facts))).tables
           for row in table.rows for cell in row.cells]
    )

    assert "flood layer hazard_flood and district polygon 1,784 km²" in text
    assert "Villages (villages_th_register)" in text and "GRP's own data" in text


def test_only_the_chosen_layers_are_listed_and_a_missing_one_says_so() -> None:
    rows = global_risk_stats({
        "selected_layers": ["schools", "early_warning_towers_ddpm"],
        "stats": {"counts": {
            "schools": {"exposed": 3, "total": 9},
            "roads": {"exposed_km": 1.0, "total_km": 5.0},
        }},
    })

    assert [row[0] for row in rows] == [
        "Schools", "Early-warning towers (early_warning_towers_ddpm)"]
    assert "returned no count" in rows[1][5]
