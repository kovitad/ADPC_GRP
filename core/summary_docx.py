"""The planner's district summary as a Word document (ADR-0033), in English or Thai.

Pure rendering: every fact arrives already gathered and screened by ``api/planning_summary.py``.
The wording follows what the Planning panel says and the rules behind it: "not exposed under
this scenario" (never "safe"), "N/A" for a centre the flood layer cannot assess, registered
village population (ADR-0027), sensitivity as a relative index (ADR-0030), no Global Risk risk
levels (ADR-0014), and Global Risk's map as an unverified link (ADR-0031).

Section 6 (ADR-0056) adds live reported flooding for Bangkok districts: the same facts as the
Planner's live answers, with camera pictures. It is never part of the assessment.

Fixed text is translated; texts that come from data, sources or Global Risk (the brief, limits,
source lines) stay as written, and the Thai document says so.
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Covers Thai and Latin; Word falls back sensibly where it is missing.
FONT = "Leelawadee UI"
INK = RGBColor(0x1C, 0x2A, 0x33)
MUTED = RGBColor(0x5B, 0x6B, 0x74)
WARN = RGBColor(0x8A, 0x53, 0x00)
TABLE_STYLE = "Table Grid"
MAX_TABLE_ROWS = 1000
LANGS = ("en", "th")

T: dict[str, dict[str, str]] = {
    "title": {"en": "Flood preparedness summary: {area}",
              "th": "สรุปการเตรียมความพร้อมรับมือน้ำท่วม: {area}"},
    "district": {"en": "District", "th": "เขต/อำเภอ"},
    "subdistrict": {"en": "Sub-district", "th": "แขวง/ตำบล"},
    "byline": {"en": "{hub} · prepared by {who} · {when} (Bangkok) · SERVIR Global Risk "
                     "Platform (GRP)",
               "th": "{hub} · จัดทำโดย {who} · {when} (เวลากรุงเทพฯ) · SERVIR Global Risk "
                     "Platform (GRP)"},
    "synthetic": {"en": "SYNTHETIC TEST DATA: invented inputs for software testing only. Do not "
                        "use for planning.",
                  "th": "ข้อมูลทดสอบสังเคราะห์: ข้อมูลสมมติสำหรับทดสอบระบบเท่านั้น ห้ามใช้วางแผน"},
    "screening": {"en": "Screening information for preparedness planning. It is not a report of "
                        "current flooding and does not certify that any place is safe.",
                  "th": "ข้อมูลคัดกรองเพื่อการวางแผนเตรียมความพร้อม ไม่ใช่รายงานสถานการณ์น้ำท่วม"
                        "ปัจจุบัน และไม่ได้รับรองว่าสถานที่ใดปลอดภัย"},
    "data_language": {"en": "", "th": "ข้อความที่มาจากแหล่งข้อมูลหรือ Global Risk (เช่น บทสรุป "
                                      "ข้อจำกัด และชื่อแหล่งข้อมูล) แสดงตามต้นฉบับภาษาอังกฤษ"},
    "s1": {"en": "1. At a glance", "th": "1. ภาพรวม"},
    "centres_in_data": {"en": "Evacuation centres in the data", "th": "ศูนย์พักพิงในข้อมูล"},
    "potentially_exposed": {"en": "Potentially exposed ({scenario})",
                            "th": "อาจได้รับผลกระทบ ({scenario})"},
    "not_exposed": {"en": "Not exposed under this scenario",
                    "th": "ไม่อยู่ในพื้นที่น้ำท่วมภายใต้สถานการณ์นี้"},
    "na_centre": {"en": "N/A (flood layer has no value at the centre)",
                  "th": "ไม่มีข้อมูล (ชั้นข้อมูลน้ำท่วมไม่มีค่าที่ตำแหน่งศูนย์)"},
    "registered_population": {"en": "Registered village population",
                              "th": "ประชากรตามทะเบียนหมู่บ้าน"},
    "people_in_extent": {"en": "Registered people in villages inside the RP{rp} modelled flood "
                               "extent",
                         "th": "ประชากรตามทะเบียนในหมู่บ้านที่อยู่ในพื้นที่น้ำท่วมจำลอง RP{rp}"},
    "item": {"en": "Item", "th": "รายการ"},
    "value": {"en": "Value", "th": "ค่า"},
    "s2": {"en": "2. Map", "th": "2. แผนที่"},
    "no_map": {"en": "No map picture was included with this download.",
               "th": "ไม่มีภาพแผนที่ในไฟล์นี้"},
    "s3": {"en": "3. Evacuation centres", "th": "3. ศูนย์พักพิง"},
    "no_centres": {"en": "The current shelter data has no records in this area. This is a gap in "
                         "the data, not a finding that it has no shelters.",
                   "th": "ข้อมูลศูนย์พักพิงปัจจุบันไม่มีรายการในพื้นที่นี้ ซึ่งเป็นช่องว่างของ"
                         "ข้อมูล ไม่ใช่ข้อสรุปว่าพื้นที่นี้ไม่มีศูนย์พักพิง"},
    "assessment_line": {"en": "Assessment {id} · {scenario} · method {method}",
                        "th": "การประเมิน {id} · {scenario} · วิธี {method}"},
    "where_move": {"en": "Where people could move", "th": "สถานที่ที่ประชาชนอาจย้ายไปได้"},
    "moveable": {"en": "{n} centre(s) are not exposed under this scenario. Check access routes, "
                       "capacity and current conditions before use.",
                 "th": "มีศูนย์ {n} แห่งที่ไม่อยู่ในพื้นที่น้ำท่วมภายใต้สถานการณ์นี้ ควรตรวจสอบ"
                       "เส้นทาง ความจุ และสภาพปัจจุบันก่อนใช้งาน"},
    "capacity_short": {"en": "capacity", "th": "ความจุ"},
    "none_moveable": {"en": "No centre in this area is identified as not exposed under this "
                            "scenario. Plan for moving people to neighbouring areas.",
                      "th": "ไม่มีศูนย์ในพื้นที่นี้ที่ระบุว่าไม่อยู่ในพื้นที่น้ำท่วมภายใต้"
                            "สถานการณ์นี้ ควรวางแผนย้ายประชาชนไปยังพื้นที่ใกล้เคียง"},
    "h_centre": {"en": "Centre", "th": "ศูนย์"},
    "h_subdistrict": {"en": "Sub-district", "th": "แขวง/ตำบล"},
    "h_village": {"en": "Village", "th": "หมู่บ้าน"},
    "h_capacity": {"en": "Capacity", "th": "ความจุ"},
    "h_status": {"en": "Status", "th": "สถานะ"},
    "h_depth": {"en": "Flood depth (m)", "th": "ความลึกน้ำ (ม.)"},
    "unknown": {"en": "unknown", "th": "ไม่ทราบ"},
    "all_centres": {"en": "All centres", "th": "ศูนย์ทั้งหมด"},
    "sensitivity_note": {"en": "Sensitivity values are the sub-district's relative index, 0 to 1. "
                               "They are not a number of people and are not used in the "
                               "assessment.",
                         "th": "ค่าความอ่อนไหวเป็นดัชนีสัมพัทธ์ของแขวง/ตำบล ระหว่าง 0 ถึง 1 "
                               "ไม่ใช่จำนวนคน และไม่ได้ใช้ในการประเมิน"},
    "s4": {"en": "4. People", "th": "4. ประชากร"},
    "people": {"en": "People", "th": "จำนวนคน"},
    "male": {"en": "Male", "th": "ชาย"},
    "female": {"en": "Female", "th": "หญิง"},
    "households": {"en": "Households", "th": "ครัวเรือน"},
    "villages_counted": {"en": "Villages with a count", "th": "หมู่บ้านที่มีข้อมูลจำนวน"},
    "of": {"en": "of", "th": "จาก"},
    "population_note": {"en": "Registered population from the village register, not a census and "
                              "not a count of vulnerable people.",
                        "th": "ประชากรตามทะเบียนหมู่บ้าน ไม่ใช่สำมะโนประชากร และไม่ใช่จำนวนกลุ่ม"
                              "เปราะบาง"},
    "no_population": {"en": "No population record for this area.",
                      "th": "ไม่มีข้อมูลประชากรของพื้นที่นี้"},
    "villages_in_extent": {"en": "Villages inside the RP{rp} modelled flood extent",
                           "th": "หมู่บ้านในพื้นที่น้ำท่วมจำลอง RP{rp}"},
    "villages": {"en": "Villages", "th": "หมู่บ้าน"},
    "registered_people": {"en": "Registered people", "th": "ประชากรตามทะเบียน"},
    "vulnerable_note": {"en": "Children, older people and people with disabilities: no headcount "
                              "exists in the data. Sensitivity maps show where they are relatively "
                              "more sensitive; the disability layer is withheld until its classes "
                              "are explained.",
                        "th": "เด็ก ผู้สูงอายุ และผู้พิการ: ไม่มีข้อมูลจำนวนคนในชุดข้อมูล แผนที่"
                              "ความอ่อนไหวแสดงเพียงพื้นที่ที่อ่อนไหวมากกว่าโดยเปรียบเทียบ และยังไม่"
                              "แสดงชั้นข้อมูลผู้พิการจนกว่าจะมีคำอธิบายระดับ"},
    "s5": {"en": "5. Global Risk evidence", "th": "5. หลักฐานจาก Global Risk"},
    "no_gr": {"en": "No Global Risk evidence has been gathered for this area in the last 30 days. "
                    "Ask about this district in Planning to add it.",
              "th": "ยังไม่มีหลักฐานจาก Global Risk สำหรับพื้นที่นี้ในช่วง 30 วันที่ผ่านมา ถามเกี่ยวกับ"
                    "พื้นที่นี้ในหน้า Planning เพื่อเพิ่มข้อมูล"},
    "gr_line": {"en": "{place} · evidence pack {pack} · gathered {when} · {status}",
                "th": "{place} · ชุดหลักฐาน {pack} · รวบรวมเมื่อ {when} · {status}"},
    "question": {"en": "Question: {q}", "th": "คำถาม: {q}"},
    "gr_flood_layer": {"en": "flood layer {h}", "th": "ชั้นข้อมูลน้ำท่วม {h}"},
    "gr_polygon": {"en": "district polygon {km2} km²", "th": "ขอบเขตพื้นที่ {km2} ตร.กม."},
    "gr_and": {"en": " and ", "th": " และ "},
    "gr_counted": {"en": "Global Risk counted with its own {what}. GRP uses the DOPA boundary and "
                         "its own flood tiles, so the two sets of counts can differ for that "
                         "reason alone.",
                   "th": "Global Risk นับโดยใช้{what}ของตนเอง ส่วน GRP ใช้ขอบเขตของกรมการปกครอง"
                         "และข้อมูลน้ำท่วมของตนเอง ตัวเลขทั้งสองจึงอาจต่างกันได้ด้วยเหตุนี้"},
    "gr_layers": {"en": "Layers chosen for this question: {layers}. Only these are listed below; "
                        "Global Risk may hold others for this district.",
                  "th": "ชั้นข้อมูลที่เลือกสำหรับคำถามนี้: {layers} แสดงเฉพาะชั้นเหล่านี้ด้านล่าง "
                        "Global Risk อาจมีชั้นข้อมูลอื่นของพื้นที่นี้"},
    "gr_adds": {"en": "What Global Risk adds for this district",
                "th": "ข้อมูลเพิ่มเติมจาก Global Risk สำหรับพื้นที่นี้"},
    "gr_adds_note": {"en": "Global Risk counts schools, hospitals, buildings and roads inside its "
                           "100-year flood hazard layer. Rows marked as GRP's own data are GRP's "
                           "datasets counted again by Global Risk, not an independent check. "
                           "\"In flood hazard\" means flood hazard class 1 or higher (1 = up to "
                           "0.5 m, 5 = over 2 m); it is not a report of current flooding.",
                     "th": "Global Risk นับโรงเรียน โรงพยาบาล อาคาร และถนนในชั้นข้อมูลภัยน้ำท่วม"
                           "รอบ 100 ปี แถวที่ระบุว่าเป็นข้อมูลของ GRP คือข้อมูลของ GRP ที่ Global "
                           "Risk นับซ้ำ ไม่ใช่การตรวจสอบอิสระ \"อยู่ในพื้นที่ภัยน้ำท่วม\" หมายถึง"
                           "ระดับ 1 ขึ้นไป (1 = ไม่เกิน 0.5 ม., 5 = เกิน 2 ม.) ไม่ใช่รายงานน้ำท่วม"
                           "ปัจจุบัน"},
    "h_in_hazard": {"en": "In flood hazard", "th": "อยู่ในพื้นที่ภัยน้ำท่วม"},
    "h_total": {"en": "Total", "th": "ทั้งหมด"},
    "h_unit": {"en": "Unit", "th": "หน่วย"},
    "h_by_class": {"en": "By depth class", "th": "ตามระดับความลึก"},
    "h_means": {"en": "What it means", "th": "ความหมาย"},
    "gr_none": {"en": "Global Risk found none of these in its flood hazard layer for this "
                      "district. That is a result, not missing data.",
                "th": "Global Risk ไม่พบรายการเหล่านี้ในชั้นข้อมูลภัยน้ำท่วมของพื้นที่นี้ ซึ่งเป็น"
                      "ผลลัพธ์ ไม่ใช่ข้อมูลที่ขาดหาย"},
    "brief": {"en": "Brief", "th": "บทสรุป"},
    "receipt": {"en": "Public receipt {id}: {url}", "th": "ใบรับรองสาธารณะ {id}: {url}"},
    "gr_map": {"en": "Global Risk map: {url}", "th": "แผนที่ Global Risk: {url}"},
    "gr_map_note": {"en": "Global Risk's own view of this receipt. GRP could not confirm which "
                          "layer its colours show (flood depth or vulnerability-weighted risk).",
                    "th": "มุมมองของ Global Risk สำหรับใบรับรองนี้ GRP ยืนยันไม่ได้ว่าสีบนแผนที่"
                          "แสดงชั้นข้อมูลใด (ความลึกน้ำ หรือความเสี่ยงถ่วงน้ำหนักความเปราะบาง)"},
    "gr_sources": {"en": "Global Risk sources", "th": "แหล่งข้อมูลของ Global Risk"},
    "h_hash": {"en": "#", "th": "#"},
    "h_title": {"en": "Title", "th": "ชื่อ"},
    "h_source": {"en": "Source", "th": "แหล่งที่มา"},
    "h_validation": {"en": "Validation", "th": "การตรวจสอบ"},
    "h_retrieval": {"en": "Retrieval", "th": "การดึงข้อมูล"},
    "gr_gaps": {"en": "Gaps Global Risk declared", "th": "ช่องว่างข้อมูลที่ Global Risk แจ้ง"},
    "s6": {"en": "6. Live reported flooding (not part of this assessment)",
           "th": "6. รายงานน้ำท่วมแบบสด (ไม่ใช่ส่วนหนึ่งของการประเมินนี้)"},
    "live_unavailable": {"en": "Live reported flooding covers Bangkok and Nonthaburi only.",
                         "th": "ข้อมูลรายงานน้ำท่วมแบบสดครอบคลุมเฉพาะกรุงเทพฯ และนนทบุรี"},
    "live_hub": {"en": "Live reported flooding is not enabled for this Hub.",
                 "th": "ยังไม่ได้เปิดข้อมูลรายงานน้ำท่วมแบบสดสำหรับ Hub นี้"},
    "live_stamp": {"en": "As of {when}, Bangkok time. Flooding reported by BMA, Traffy Fondue and "
                         "the public through Floodboard, grouped by GRP. Not a flood map, not a "
                         "warning, and not part of the assessment above.",
                   "th": "ข้อมูล ณ {when} เวลากรุงเทพฯ น้ำท่วมที่รายงานโดย กทม. Traffy Fondue และ"
                         "ประชาชนผ่าน Floodboard จัดกลุ่มโดย GRP ไม่ใช่แผนที่น้ำท่วม ไม่ใช่การเตือนภัย "
                         "และไม่ใช่ส่วนหนึ่งของการประเมินข้างต้น"},
    "live_stale": {"en": "The latest snapshot is {minutes} minutes old; read it as recent "
                         "history, not as now.",
                   "th": "ข้อมูลล่าสุดมีอายุ {minutes} นาที ควรอ่านเป็นประวัติล่าสุด ไม่ใช่สถานการณ์"
                         "ขณะนี้"},
    "live_rollup": {"en": "Shown for the whole of {district}, not only the selected sub-district.",
                    "th": "แสดงข้อมูลทั้ง{district} ไม่ใช่เฉพาะแขวงที่เลือก"},
    "live_glance": {"en": "At a glance", "th": "ภาพรวม"},
    "live_active": {"en": "Active incidents", "th": "เหตุการณ์ที่ยังมีรายงาน"},
    "live_conf": {"en": "  {word}", "th": "  {word}"},
    "live_changes": {"en": "Last {w} minutes: new / grew / no longer reported",
                     "th": "{w} นาทีที่ผ่านมา: ใหม่ / ขยายตัว / ไม่มีรายงานแล้ว"},
    "live_fac_count": {"en": "Facilities with flooding reported nearby",
                       "th": "สถานที่สำคัญที่มีรายงานน้ำท่วมใกล้เคียง"},
    "live_ddpm_count": {"en": "  of which DDPM evacuation centres",
                        "th": "  ในจำนวนนี้เป็นศูนย์พักพิงของ ปภ."},
    "live_incidents": {"en": "Incidents", "th": "เหตุการณ์"},
    "h_roads": {"en": "Roads", "th": "ถนน"},
    "h_confidence": {"en": "Confidence", "th": "ความมั่นใจ"},
    "h_reports": {"en": "Reports", "th": "รายงาน"},
    "h_deepest": {"en": "Deepest reported (cm)", "th": "ลึกสุดที่รายงาน (ซม.)"},
    "h_last": {"en": "Last report", "th": "รายงานล่าสุด"},
    "live_more": {"en": "{n} more incident(s) are not listed.",
                  "th": "มีอีก {n} เหตุการณ์ที่ไม่ได้แสดง"},
    "live_none": {"en": "No incident with flooding reported is open in this district now.",
                  "th": "ขณะนี้ไม่มีเหตุการณ์ที่มีรายงานน้ำท่วมในเขตนี้"},
    "live_facilities": {"en": "Facilities with flooding reported nearby",
                        "th": "สถานที่สำคัญที่มีรายงานน้ำท่วมใกล้เคียง"},
    "h_name": {"en": "Name", "th": "ชื่อ"},
    "h_type": {"en": "Type", "th": "ประเภท"},
    "h_distance": {"en": "Flooded road within (m)", "th": "ถนนที่มีรายงานน้ำท่วมห่าง (ม.)"},
    "live_fac_note": {"en": "\"Flooding reported nearby\" means a road with current reports lies "
                            "close by. It does not mean the building is flooded, and access is "
                            "not confirmed.",
                      "th": "\"มีรายงานน้ำท่วมใกล้เคียง\" หมายถึงมีถนนที่มีรายงานน้ำท่วมอยู่ใกล้ "
                            "ไม่ได้หมายความว่าอาคารถูกน้ำท่วม และยังไม่ได้ยืนยันการเข้าถึง"},
    "live_rain": {"en": "Rain now and in 30 minutes (radar {when})",
                  "th": "ฝนขณะนี้และใน 30 นาที (เรดาร์ {when})"},
    "h_district": {"en": "District", "th": "เขต"},
    "h_rain_now": {"en": "Now", "th": "ขณะนี้"},
    "h_heaviest": {"en": "Heaviest now", "th": "หนักสุดขณะนี้"},
    "h_in30": {"en": "In 30 minutes", "th": "ใน 30 นาที"},
    "live_rain_note": {"en": "Rain is not flooding, and a forecast is not an observation.",
                       "th": "ฝนไม่ใช่น้ำท่วม และการพยากรณ์ไม่ใช่การสังเกตการณ์"},
    "live_cameras": {"en": "Cameras near the listed incidents", "th": "กล้องใกล้เหตุการณ์ที่แสดง"},
    "live_camera_caption": {"en": "{incident} · {name} · {distance} m from the reported flooding · "
                                  "{credit}, {when}",
                            "th": "{incident} · {name} · ห่างจากจุดที่มีรายงาน {distance} ม. · "
                                  "{credit} เวลา {when}"},
    "live_camera_note": {"en": "Camera pictures show one moment from one angle. An empty road in "
                               "a picture is not proof that the area is dry.",
                         "th": "ภาพจากกล้องแสดงเพียงช่วงเวลาและมุมเดียว ถนนที่ดูแห้งในภาพไม่ใช่"
                               "หลักฐานว่าพื้นที่ไม่มีน้ำท่วม"},
    "h_camera": {"en": "Camera", "th": "กล้อง"},
    "h_incident": {"en": "Incident", "th": "เหตุการณ์"},
    "h_live_view": {"en": "Live view", "th": "ดูภาพสด"},
    "no_picture": {"en": "no picture", "th": "ไม่มีภาพ"},
    "live_credit": {"en": "Sources: {credit}.", "th": "แหล่งข้อมูล: {credit}"},
    "s7": {"en": "7. What this cannot tell you", "th": "7. ข้อจำกัดของเอกสารนี้"},
    "s8": {"en": "8. Sources and versions", "th": "8. แหล่งข้อมูลและรุ่นข้อมูล"},
    "h_role": {"en": "Role", "th": "บทบาท"},
    "h_provider": {"en": "Provider", "th": "ผู้ให้ข้อมูล"},
    "h_version": {"en": "Version", "th": "รุ่น"},
    "generated_note": {"en": "Generated by GRP from the data versions above. Numbers are copied "
                             "from the stored records; nothing in this document was estimated "
                             "for it.",
                       "th": "สร้างโดย GRP จากข้อมูลรุ่นข้างต้น ตัวเลขคัดลอกจากข้อมูลที่จัดเก็บ "
                             "ไม่มีการประมาณค่าเพิ่มสำหรับเอกสารนี้"},
    "rows_shown": {"en": "Showing the first {shown} of {total} rows; the CSV download has them "
                         "all.",
                   "th": "แสดง {shown} แถวแรกจาก {total} แถว ดาวน์โหลด CSV เพื่อดูทั้งหมด"},
    "not_recorded": {"en": "not recorded", "th": "ไม่มีข้อมูล"},
}

STATUS_LABELS = {
    "potentially_exposed": {"en": "Potentially exposed", "th": "อาจได้รับผลกระทบ"},
    "not_exposed_under_scenario": {"en": "Not exposed under this scenario",
                                   "th": "ไม่อยู่ในพื้นที่น้ำท่วมภายใต้สถานการณ์นี้"},
    "unable_to_assess": {"en": "N/A", "th": "ไม่มีข้อมูล"},
    "not_assessed": {"en": "Not assessed", "th": "ยังไม่ได้ประเมิน"},
}
CONFIDENCE = {
    "high": {"en": "High confidence", "th": "ความมั่นใจสูง"},
    "medium": {"en": "Medium confidence", "th": "ความมั่นใจปานกลาง"},
    "low": {"en": "Low confidence", "th": "ความมั่นใจต่ำ"},
    "conflicting": {"en": "Conflicting evidence", "th": "หลักฐานขัดแย้ง"},
}
FACILITY_TYPES = {
    "evacuation_centre": {"en": "DDPM evacuation centre", "th": "ศูนย์พักพิง ปภ."},
    "school": {"en": "School", "th": "โรงเรียน"},
    "hospital": {"en": "Hospital", "th": "โรงพยาบาล"},
    "clinic": {"en": "Clinic", "th": "คลินิก"},
}


def t(key: str, lang: str, **values: Any) -> str:
    text = T[key]["th" if lang == "th" else "en"]
    return text.format(**values) if values else text


def _local(value: Any, lang: str) -> str:
    """A text that may come in both languages, such as a camera name."""

    if isinstance(value, dict):
        # The other language when this one is blank or only punctuation (such as "- · -").
        order = ("th", "en") if lang == "th" else ("en", "th")
        for code in order:
            text = str(value.get(code) or "")
            if re.search(r"\w", text):
                return text
        return ""
    return "" if value is None else str(value)


def _label(table: dict[str, dict[str, str]], key: Any, lang: str) -> str:
    entry = table.get(str(key))
    return entry["th" if lang == "th" else "en"] if entry else str(key)


def _set_font(style_or_run: Any, size: float | None = None) -> None:
    font = style_or_run.font
    font.name = FONT
    if size:
        font.size = Pt(size)
    element = style_or_run.element
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    # w:cs is the slot Word uses for Thai; without it Thai names can fall back to boxes.
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(slot), FONT)


def _styles(document: Any) -> None:
    for name, size in (("Normal", 10.5), ("Heading 1", 16), ("Heading 2", 13),
                       ("Heading 3", 11.5), ("Title", 22), ("Caption", 9)):
        style = document.styles[name]
        _set_font(style, size)
    document.styles["Normal"].font.color.rgb = INK
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)


def _para(document: Any, text: str, *, bold: bool = False, color: RGBColor | None = None,
          size: float | None = None, style: str | None = None) -> Any:
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return paragraph


def _note(document: Any, text: str) -> None:
    _para(document, text, color=MUTED, size=9)


def _warn(document: Any, text: str) -> None:
    _para(document, text, color=WARN, size=9.5)


def _bullets(document: Any, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(str(item), style="List Bullet")


def _table(document: Any, header: list[str], rows: list[list[Any]], lang: str = "en") -> None:
    table = document.add_table(rows=1, cols=len(header))
    table.style = TABLE_STYLE
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows[:MAX_TABLE_ROWS]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = ""
            run = cell.paragraphs[0].add_run("" if value is None else str(value))
            run.font.size = Pt(9)
    if len(rows) > MAX_TABLE_ROWS:
        _note(document, t("rows_shown", lang, shown=f"{MAX_TABLE_ROWS:,}", total=f"{len(rows):,}"))


def _number(value: Any, lang: str = "en") -> str:
    if value is None:
        return t("not_recorded", lang)
    if isinstance(value, float):
        return f"{value:,.1f}"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def _area_name(area: dict[str, Any], lang: str = "en") -> str:
    level = t("district" if area.get("admin_level") == "district" else "subdistrict", lang)
    if lang == "th" and area.get("name_th"):
        return f"{level}{area['name_th']} ({area.get('name', '')})"
    name = f"{area.get('name', '')} {level}".strip()
    if area.get("name_th"):
        name += f" ({area['name_th']})"
    return name


def _live_section(document: Any, live: dict[str, Any] | None, lang: str) -> None:
    document.add_heading(t("s6", lang), level=1)
    if not live or not live.get("available"):
        reason = (live or {}).get("reason")
        _note(document, t("live_hub" if reason == "hub_not_enabled" else "live_unavailable", lang))
        return
    _warn(document, t("live_stamp", lang, when=live["as_of"]))
    if live.get("stale"):
        _warn(document, t("live_stale", lang, minutes=live["age_minutes"]))
    if live.get("rolled_up_from"):
        district = (live.get("district_name_th") if lang == "th" else None) or live["district_name"]
        _note(document, t("live_rollup", lang, district=district))

    situation, changes = live["situation"], live["changes"]
    facilities = live.get("facilities") or []
    document.add_heading(t("live_glance", lang), level=2)
    rows = [[t("live_active", lang), _number(situation.get("active_incidents"), lang)]]
    for word, count in (situation.get("incidents_by_confidence") or {}).items():
        rows.append([t("live_conf", lang, word=_label(CONFIDENCE, word, lang)), _number(count)])
    rows.append([t("live_changes", lang, w=changes.get("window_minutes")),
                 f"{changes.get('new_incidents', 0)} / {changes.get('incidents_grew', 0)} / "
                 f"{changes.get('no_longer_reported', 0)}"])
    rows.append([t("live_fac_count", lang), _number(len(facilities))])
    rows.append([t("live_ddpm_count", lang), _number(
        sum(1 for f in facilities if f.get("facility_type") == "evacuation_centre"))])
    _table(document, [t("item", lang), t("value", lang)], rows, lang)

    document.add_heading(t("live_incidents", lang), level=2)
    incidents = live.get("incidents") or []
    if incidents:
        _table(document, ["#", t("h_roads", lang), t("h_confidence", lang), t("h_reports", lang),
                          t("h_deepest", lang), t("h_last", lang)], [
            [i["id"], ", ".join(i.get("roads") or []),
             _label(CONFIDENCE, i.get("confidence"), lang), i.get("reports"),
             "" if i.get("deepest_reported_cm") is None else i["deepest_reported_cm"],
             i.get("newest_evidence_bangkok")]
            for i in incidents
        ], lang)
        if live.get("incidents_not_listed"):
            _note(document, t("live_more", lang, n=live["incidents_not_listed"]))
    else:
        _note(document, t("live_none", lang))

    if facilities:
        document.add_heading(t("live_facilities", lang), level=2)
        _table(document, [t("h_name", lang), t("h_type", lang), t("h_distance", lang),
                          t("h_source", lang)], [
            [f.get("name"), _label(FACILITY_TYPES, f.get("facility_type"), lang),
             f.get("flooding_reported_within_m"), f.get("source")]
            for f in facilities
        ], lang)
        _note(document, t("live_fac_note", lang))

    rain = live.get("rain")
    if rain and rain.get("districts"):
        document.add_heading(t("live_rain", lang, when=rain.get("radar_time_bangkok")), level=2)
        _table(document, [t("h_district", lang), t("h_rain_now", lang), t("h_heaviest", lang),
                          t("h_in30", lang)],
               [[d.get("district"), d.get("rain_now"), d.get("heaviest_now"), d.get("in_30_min")]
                for d in rain["districts"]], lang)
        _note(document, t("live_rain_note", lang))

    cameras = live.get("cameras") or []
    if cameras:
        document.add_heading(t("live_cameras", lang), level=2)
        for camera in cameras:
            if not camera.get("picture"):
                continue
            try:
                document.add_picture(BytesIO(camera["picture"]), width=Cm(12))
            except Exception:  # an unreadable picture is listed below instead
                camera["picture"] = None
                continue
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _para(document, t("live_camera_caption", lang, incident=camera["incident"],
                              name=_local(camera.get("name"), lang),
                              distance=camera["distance_m"],
                              credit=camera.get("credit") or camera.get("source") or "",
                              when=camera.get("picture_at") or ""), style="Caption")
        header = [t("h_incident", lang), t("h_camera", lang), t("h_distance", lang),
                  t("h_source", lang)]
        rows = [[c["incident"], _local(c.get("name"), lang), c["distance_m"],
                 _local(c.get("source"), lang)] for c in cameras]
        if any(c.get("viewer_url") for c in cameras):
            header.append(t("h_live_view", lang))
            for row, c in zip(rows, cameras, strict=True):
                row.append(c.get("viewer_url") or "")
        _table(document, header, rows, lang)
        _note(document, t("live_camera_note", lang))

    _note(document, t("live_credit", lang, credit=live.get("credit") or ""))
    no_warnings = live.get("no_warnings") or {}
    for code in ("th", "en") if lang == "th" else ("en", "th"):
        if no_warnings.get(code):
            _para(document, no_warnings[code], bold=True, size=9.5)


def render_summary(facts: dict[str, Any], lang: str = "en") -> bytes:
    lang = lang if lang in LANGS else "en"
    document = Document()
    _styles(document)
    area = facts["area"]
    province = area.get("province_name") or ""
    if area.get("province_name_th"):
        province = (f"{area['province_name_th']} ({province})" if lang == "th"
                    else f"{province} ({area['province_name_th']})")

    document.add_paragraph(t("title", lang, area=_area_name(area, lang)), style="Title")
    _para(document, ", ".join(p for p in (province, area.get("country_name") or "") if p),
          size=12)
    _note(document, t("byline", lang, hub=facts["hub"], who=facts["prepared_by"],
                      when=facts["generated_at"]))
    if facts.get("synthetic"):
        _warn(document, t("synthetic", lang))
    _warn(document, t("screening", lang))
    if lang == "th":
        _note(document, t("data_language", lang))

    # 1. At a glance
    document.add_heading(t("s1", lang), level=1)
    centres = facts["centres"]
    glance: list[list[Any]] = [[t("centres_in_data", lang), _number(centres["total"], lang)]]
    summary = centres.get("summary")
    if summary:
        glance += [
            [t("potentially_exposed", lang, scenario=centres["scenario"]),
             _number(summary.get("potentially_exposed"), lang)],
            [t("not_exposed", lang), _number(summary.get("not_exposed_under_scenario"), lang)],
            [t("na_centre", lang), _number(summary.get("unable_to_assess"), lang)],
        ]
    people = facts.get("people") or {}
    population = people.get("population")
    if population:
        glance.append([t("registered_population", lang),
                       _number(population.get("total_population"), lang)])
    exposure = people.get("flood_exposure")
    if exposure:
        glance.append([t("people_in_extent", lang, rp=exposure.get("return_period_years")),
                       _number(exposure.get("people_in_zone"), lang)])
    for item in facts.get("supporting") or []:
        glance.append([item["title"], _number(item["count"], lang)])
    _table(document, [t("item", lang), t("value", lang)], glance, lang)

    # 2. Map
    document.add_heading(t("s2", lang), level=1)
    picture = (facts.get("map") or {}).get("png")
    if picture:
        document.add_picture(BytesIO(picture), width=Cm(16.5))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        _note(document, t("no_map", lang))
    _para(document, (facts.get("map") or {}).get("caption") or "", style="Caption")

    # 3. Evacuation centres
    document.add_heading(t("s3", lang), level=1)
    _note(document, centres["source_line"])
    if centres["total"] == 0:
        _warn(document, (centres.get("gap_note") if lang == "en" else None)
              or t("no_centres", lang))
    if summary:
        _para(document, t("assessment_line", lang, id=centres["assessment_id"],
                          scenario=centres["scenario"], method=centres["method"]), size=9.5)
    moveable = [row for row in centres["rows"]
                if row.get("status") == "not_exposed_under_scenario"]
    if summary:
        document.add_heading(t("where_move", lang), level=2)
        if moveable:
            _para(document, t("moveable", lang, n=f"{len(moveable):,}"))
            _bullets(document, [
                f"{row['name']}" + (f" · {row['subdistrict']}" if row.get("subdistrict") else "")
                + (f" · {t('capacity_short', lang)} {row['capacity']}" if row.get("capacity")
                   else "")
                for row in moveable[:60]
            ])
        else:
            _para(document, t("none_moveable", lang))
    indicator_titles = centres.get("indicator_titles") or {}
    header = [t("h_centre", lang), t("h_subdistrict", lang), t("h_village", lang),
              t("h_capacity", lang), t("h_status", lang), t("h_depth", lang)]
    header += list(indicator_titles.values())
    rows = []
    for row in centres["rows"]:
        values = [
            row.get("name"), row.get("subdistrict"), row.get("village"),
            row.get("capacity") if row.get("capacity") not in (None, "") else t("unknown", lang),
            _label(STATUS_LABELS, row.get("status"), lang),
            "" if row.get("depth_m") is None else f"{row['depth_m']:.2f}",
        ]
        for key in indicator_titles:
            value = (row.get("indicators") or {}).get(key)
            values.append("" if value is None else f"{value:.2f}")
        rows.append(values)
    if rows:
        document.add_heading(t("all_centres", lang), level=2)
        _table(document, header, rows, lang)
    if indicator_titles:
        _note(document, t("sensitivity_note", lang))

    # 4. People
    document.add_heading(t("s4", lang), level=1)
    if population:
        _table(document, [t("registered_population", lang), t("value", lang)], [
            [t("people", lang), _number(population.get("total_population"), lang)],
            [t("male", lang), _number(population.get("male"), lang)],
            [t("female", lang), _number(population.get("female"), lang)],
            [t("households", lang), _number(population.get("households"), lang)],
            [t("villages_counted", lang),
             f"{_number(population.get('counted_village_count'), lang)} {t('of', lang)} "
             f"{_number(population.get('village_count'), lang)}"],
        ], lang)
        _note(document, t("population_note", lang))
    else:
        _note(document, t("no_population", lang))
    if exposure:
        document.add_heading(t("villages_in_extent", lang, rp=exposure.get("return_period_years")),
                             level=2)
        _table(document, [t("item", lang), t("value", lang)], [
            [t("villages", lang), _number(exposure.get("villages_in_zone"), lang)],
            [t("registered_people", lang), _number(exposure.get("people_in_zone"), lang)],
            [t("households", lang), _number(exposure.get("households_in_zone"), lang)],
        ], lang)
        if exposure.get("caveat"):
            _note(document, exposure["caveat"])
    _note(document, t("vulnerable_note", lang))

    # 5. Global Risk evidence
    document.add_heading(t("s5", lang), level=1)
    evidence = facts.get("global_risk")
    if not evidence:
        _note(document, t("no_gr", lang))
    else:
        _para(document, t("gr_line", lang, place=evidence["place"], pack=evidence["pack_id"],
                          when=evidence["assembled_at"], status=evidence["status"]), size=9.5)
        if evidence.get("question"):
            _note(document, t("question", lang, q=evidence["question"]))
        if evidence.get("source_note"):
            _warn(document, evidence["source_note"])
        if evidence.get("hazard") or evidence.get("area_km2") is not None:
            counted = [t("gr_flood_layer", lang, h=evidence["hazard"])] if evidence.get(
                "hazard") else []
            if evidence.get("area_km2") is not None:
                counted.append(t("gr_polygon", lang, km2=f"{evidence['area_km2']:,}"))
            _note(document, t("gr_counted", lang, what=t("gr_and", lang).join(counted)))
        if evidence.get("selected_layers"):
            _note(document, t("gr_layers", lang, layers=", ".join(evidence["selected_layers"])))
        if evidence.get("stats"):
            document.add_heading(t("gr_adds", lang), level=2)
            _para(document, t("gr_adds_note", lang), size=9.5)
            _table(document, [t("item", lang), t("h_in_hazard", lang), t("h_total", lang),
                              t("h_unit", lang), t("h_by_class", lang), t("h_means", lang)],
                   evidence["stats"], lang)
            if all(row[1] == 0 for row in evidence["stats"]):
                _note(document, t("gr_none", lang))
        if evidence.get("brief"):
            document.add_heading(t("brief", lang), level=2)
            for line in str(evidence["brief"]).splitlines():
                text = line.strip()
                if not text:
                    continue
                if text.startswith("#"):
                    document.add_heading(text.lstrip("#").strip(), level=3)
                elif text[:2] in {"- ", "* "}:
                    document.add_paragraph(text[2:].replace("**", ""), style="List Bullet")
                else:
                    document.add_paragraph(text.replace("**", ""))
        if evidence.get("receipt_id"):
            _para(document, t("receipt", lang, id=evidence["receipt_id"],
                              url=evidence.get("receipt_url") or ""), size=9.5)
        if evidence.get("map_link"):
            _para(document, t("gr_map", lang, url=evidence["map_link"]), size=9.5)
            _note(document, t("gr_map_note", lang))
        if evidence.get("citations"):
            document.add_heading(t("gr_sources", lang), level=2)
            _table(document, [t("h_hash", lang), t("h_title", lang), t("h_source", lang),
                              t("h_validation", lang), t("h_retrieval", lang)],
                   evidence["citations"], lang)
        if evidence.get("gaps"):
            document.add_heading(t("gr_gaps", lang), level=2)
            _bullets(document, evidence["gaps"])

    # 6. Live reported flooding (ADR-0056)
    _live_section(document, facts.get("live"), lang)

    # 7. Limits
    document.add_heading(t("s7", lang), level=1)
    _bullets(document, facts.get("limits") or [])

    # 8. Sources
    document.add_heading(t("s8", lang), level=1)
    _table(document, [t("h_role", lang), t("h_title", lang), t("h_provider", lang),
                      t("h_version", lang)], facts.get("sources") or [], lang)
    _note(document, t("generated_note", lang))

    out = BytesIO()
    document.save(out)
    return out.getvalue()
